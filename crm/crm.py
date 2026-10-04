#!/usr/bin/env python3
"""Мини-CRM брокера: база контактов из телефонной книги, категории (теги)
и адресные рассылки предложений по недвижимости в Telegram, SMS и WhatsApp
с вашего личного номера. Вся история отправок хранится в базе.

Быстрый старт — см. README.md рядом. Коротко:
    python crm.py import contacts.vcf        # контакты из телефона
    python crm.py export contacts.csv        # разметить категории в Excel
    python crm.py import contacts.csv        # загрузить разметку обратно
    python crm.py send --name "Долина" --template templates/example_offer.txt \
        --tags инвестор --channels tg,wa,sms            # пробный прогон
    ... тот же вызов с --send                           # реальная отправка
"""

import argparse
import configparser
import csv
import datetime as dt
import html
import io
import quopri
import random
import re
import sqlite3
import sys
import time
import urllib.parse
from pathlib import Path

import classify

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "crm.db"
CONFIG_PATH = BASE_DIR / "config.ini"
RULES_PATH = BASE_DIR / "categories.txt"
CHANNELS = ("tg", "sms", "wa")

SCHEMA = """
CREATE TABLE IF NOT EXISTS contacts (
    id         INTEGER PRIMARY KEY,
    name       TEXT NOT NULL DEFAULT '',
    call_name  TEXT NOT NULL DEFAULT '',
    phone      TEXT UNIQUE,
    telegram   TEXT NOT NULL DEFAULT '',
    tags       TEXT NOT NULL DEFAULT '',
    consent    INTEGER NOT NULL DEFAULT 0,
    stopped    INTEGER NOT NULL DEFAULT 0,
    notes      TEXT NOT NULL DEFAULT '',
    gender     TEXT NOT NULL DEFAULT '',
    region     TEXT NOT NULL DEFAULT '',
    operator   TEXT NOT NULL DEFAULT '',
    sphere     TEXT NOT NULL DEFAULT '',
    in_phonebook INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS campaigns (
    id         INTEGER PRIMARY KEY,
    name       TEXT UNIQUE NOT NULL,
    template   TEXT NOT NULL,
    photo      TEXT NOT NULL DEFAULT '',
    tags       TEXT NOT NULL DEFAULT '',
    channels   TEXT NOT NULL,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS messages (
    id          INTEGER PRIMARY KEY,
    campaign_id INTEGER NOT NULL REFERENCES campaigns(id),
    contact_id  INTEGER NOT NULL REFERENCES contacts(id),
    channel     TEXT NOT NULL,
    status      TEXT NOT NULL,
    error       TEXT NOT NULL DEFAULT '',
    text        TEXT NOT NULL DEFAULT '',
    sent_at     TEXT NOT NULL,
    UNIQUE (campaign_id, contact_id, channel)
);
-- История взаимодействий: сообщения (в обе стороны), звонки, встречи,
-- заметки, откуда человек попал в базу.
CREATE TABLE IF NOT EXISTS interactions (
    id          INTEGER PRIMARY KEY,
    contact_id  INTEGER NOT NULL REFERENCES contacts(id),
    at          TEXT NOT NULL,
    channel     TEXT NOT NULL DEFAULT '',
    direction   TEXT NOT NULL DEFAULT '',
    kind        TEXT NOT NULL DEFAULT '',
    text        TEXT NOT NULL DEFAULT '',
    status      TEXT NOT NULL DEFAULT '',
    ext_id      TEXT,
    UNIQUE (contact_id, ext_id)
);
CREATE INDEX IF NOT EXISTS interactions_contact ON interactions (contact_id, at);
"""

CHANNEL_TITLES = {"tg": "Telegram", "wa": "WhatsApp", "sms": "SMS", "call": "Звонок",
                  "meet": "Встреча", "email": "Email", "": ""}

# Статусы, при которых человек считается охваченным кампанией.
DELIVERED = ("sent", "queued", "link")


def now():
    return dt.datetime.now().isoformat(timespec="seconds")


def connect(path=None):
    path = Path(path or DB_PATH)
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    db.executescript(SCHEMA)
    # База, созданная прошлой версией: добавляем новые колонки.
    cols = {r["name"] for r in db.execute("PRAGMA table_info(contacts)")}
    for col in ("call_name", "gender", "region", "operator", "sphere"):
        if col not in cols:
            db.execute(f"ALTER TABLE contacts ADD COLUMN {col} TEXT NOT NULL DEFAULT ''")
    if "in_phonebook" not in cols:
        db.execute("ALTER TABLE contacts ADD COLUMN in_phonebook INTEGER NOT NULL DEFAULT 0")
        # Всё, что не пришло из списков чатов, загружено из телефона.
        db.execute("UPDATE contacts SET in_phonebook=1 WHERE notes NOT LIKE 'из чата%'")
    db.commit()
    return db


def log_interaction(db, contact_id, kind, text="", channel="", direction="",
                    status="", at=None, ext_id=None):
    db.execute(
        "INSERT OR IGNORE INTO interactions (contact_id, at, channel, direction,"
        " kind, text, status, ext_id) VALUES (?,?,?,?,?,?,?,?)",
        (contact_id, at or now(), channel, direction, kind, text, status, ext_id))


def timeline(db, contact_id):
    """Вся история человека, от новых к старым: взаимодействия и рассылки."""
    items = [dict(r) for r in db.execute(
        "SELECT at, channel, direction, kind, text, status FROM interactions"
        " WHERE contact_id=?", (contact_id,))]
    for m in db.execute(
            "SELECT m.sent_at AS at, m.channel, m.status, m.text, c.name AS campaign"
            " FROM messages m JOIN campaigns c ON c.id=m.campaign_id"
            " WHERE m.contact_id=?", (contact_id,)):
        items.append({"at": m["at"], "channel": m["channel"], "direction": "out",
                      "kind": f"рассылка «{m['campaign']}»", "text": m["text"],
                      "status": m["status"]})
    return sorted(items, key=lambda x: x["at"], reverse=True)


def telegram_configured(cfg):
    """Заполнен ли [telegram] в config.ini (а не образец из config.example.ini)."""
    if not cfg.has_section("telegram"):
        return False
    t = cfg["telegram"]
    return t.get("api_id", "").strip().isdigit() and t.get("api_id") != "1234567" \
        and len(t.get("api_hash", "").strip()) == 32 \
        and not t.get("api_hash", "").startswith("0123456789abcdef")


def load_config():
    cfg = configparser.ConfigParser()
    if CONFIG_PATH.exists():
        cfg.read(CONFIG_PATH, encoding="utf-8")
    return cfg


# ---------------------------------------------------------------- телефоны, теги

def normalize_phone(raw):
    """Приводит номер к виду +7XXXXXXXXXX. Возвращает None, если это не номер."""
    if not raw:
        return None
    raw = raw.strip()
    digits = re.sub(r"\D", "", raw)
    if raw.startswith("+"):
        if digits.startswith("7"):  # Россия/Казахстан: ровно 11 цифр
            return "+" + digits if len(digits) == 11 else None
        return "+" + digits if 10 <= len(digits) <= 15 else None
    if len(digits) == 11 and digits[0] in "78":
        return "+7" + digits[1:]
    if len(digits) == 10 and digits[0] == "9":
        return "+7" + digits
    if 11 < len(digits) <= 15:
        return "+" + digits
    return None


def phone_ok(phone):
    return bool(phone) and normalize_phone(phone) == phone


# Слова в имени контакта, которые обозначают не имя, а пометку.
NOT_A_NAME = {"покупатель", "продавец", "собственник", "клиент", "аренда",
              "арендатор", "инвестор", "риелтор", "риэлтор", "агент", "агентство",
              "застройщик", "ооо", "ип", "неотвеченный", "мама", "папа"}


def guess_first_name(name):
    """Имя для обращения: «. Светлана Владимировна» -> «Светлана»,
    «2ой Покупатель Снт» -> '' (лучше без имени, чем с ошибкой)."""
    known = classify.find_first_name(name)
    if known:
        return known
    words = [w for w in (name or "").split() if re.search(r"\w", w)]
    if not words:
        return ""
    first = words[0].strip(".,-_()")
    if not re.fullmatch(r"[A-Za-zА-Яа-яЁё]+(-[A-Za-zА-Яа-яЁё]+)?", first):
        return ""
    if first.lower() in NOT_A_NAME or len(first) < 2:
        return ""
    return first[0].upper() + first[1:]


def split_tags(value):
    return [t.strip().lower() for t in re.split(r"[;,]", value or "") if t.strip()]


def join_tags(tags):
    seen = []
    for t in tags:
        if t and t not in seen:
            seen.append(t)
    return ",".join(seen)


def normalize_telegram(value):
    value = (value or "").strip()
    value = re.sub(r"^(https?://)?t\.me/", "", value)
    return value.lstrip("@")


# ---------------------------------------------------------------- импорт

def parse_vcf(text):
    """Разбирает .vcf (экспорт контактов iPhone / Android / Google)."""
    # Склеиваем перенесённые строки (RFC 6350) и мягкие переносы quoted-printable.
    text = text.replace("\r\n", "\n")
    text = re.sub(r"\n[ \t]", "", text)
    text = re.sub(r"=\n", "", text)

    contacts, cur = [], None
    for line in text.split("\n"):
        if not line.strip():
            continue
        upper = line.upper()
        if upper.startswith("BEGIN:VCARD"):
            cur = {"name": "", "n": "", "org": "", "phones": [], "tags": [],
                   "notes": ""}
            continue
        if upper.startswith("END:VCARD"):
            if cur is not None:
                cur["name"] = cur["name"] or cur["n"] or cur["org"]
                # Поле «Компания»: там часто стоят пометки вроде «Wlc», «агент».
                if cur["org"] and cur["org"].lower() not in cur["name"].lower():
                    cur["name"] = f"{cur['name']} {cur['org']}".strip()
                del cur["n"], cur["org"]
                contacts.append(cur)
            cur = None
            continue
        if cur is None or ":" not in line:
            continue
        head, value = line.split(":", 1)
        parts = head.split(";")
        key = parts[0].split(".")[-1].upper()  # "item1.TEL" -> "TEL"
        params = [p.upper() for p in parts[1:]]
        if any("QUOTED-PRINTABLE" in p for p in params):
            charset = next((p.split("=", 1)[1] for p in params
                            if p.startswith("CHARSET=")), "UTF-8")
            value = quopri.decodestring(value.encode("ascii", "ignore")) \
                .decode(charset, "replace")
        value = value.replace("\\,", ",").replace("\\;", ";").replace("\\n", " ")

        if key == "FN":
            cur["name"] = value.strip()
        elif key == "N":
            last, first = (value.split(";") + ["", ""])[:2]
            cur["n"] = " ".join(x for x in (first, last) if x).strip()
        elif key == "ORG":
            cur["org"] = " ".join(x.strip() for x in value.split(";") if x.strip())
        elif key == "TEL":
            cur["phones"].append(value)
        elif key == "CATEGORIES":
            cur["tags"] += split_tags(value)
        elif key == "NOTE":
            cur["notes"] = value.strip()
    return contacts


def read_csv_text(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1251"):  # Excel в Windows сохраняет в cp1251
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


# Колонки таблицы export/import: (ключ, заголовок).
COLUMNS = [
    ("id", "id"),
    ("phone", "Телефон"),
    ("name", "Имя в телефоне"),
    ("call_name", "Обращение"),
    ("gender", "Пол"),
    ("region", "Регион номера"),
    ("operator", "Оператор"),
    ("sphere", "Сфера деятельности"),
    ("categories", "Категории"),
    ("chats", "Списки (чаты, мероприятия, сделки)"),
    ("sources", "Где есть"),
    ("overlap", "Пересечение"),
    ("consent", "Согласие"),
    ("stopped", "СТОП"),
    ("telegram", "Telegram"),
    ("notes", "Заметки"),
]
HEADER_TO_KEY = {h.lower(): k for k, h in COLUMNS}
HEADER_TO_KEY.update({k: k for k, _ in COLUMNS})
HEADER_TO_KEY["tags"] = "categories"  # файлы прошлой версии
HEADER_TO_KEY["чаты"] = "chats"

CHAT_PREFIX = "чат-"
# Пометки-источники: откуда человек попал в базу, кроме телефона.
SOURCE_PREFIXES = (CHAT_PREFIX, "мероприятие-", "сделки-")
CHAT_TITLES = {"чат-wlc": "WLC", "чат-premium-wlc": "Premium WLC"}


def is_source(tag):
    return tag.startswith(SOURCE_PREFIXES)


def chat_title(tag):
    return CHAT_TITLES.get(tag, tag)


def chat_tag(title):
    """«Premium WLC» -> «чат-premium-wlc»; «мероприятие-…» остаётся как есть."""
    t = title.strip().lower()
    for tag, name in CHAT_TITLES.items():
        if t == name.lower():
            return tag
    return t if is_source(t) else CHAT_PREFIX + t.replace(" ", "-")


def yes_no(value):
    v = str(value or "").strip().lower()
    if v in ("1", "да", "yes", "y", "+", "true", "1.0"):
        return 1
    if v in ("0", "нет", "no", "n", "-", "false", "0.0"):
        return 0
    return ""


def read_grid(path):
    """Все листы файла как [(лист, [[ячейки строки], ...]), ...]."""
    if str(path).lower().endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError:
            sys.exit("Для .xlsx установите: python3 -m pip install --user openpyxl")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = [["" if v is None else str(v).strip() for v in r]
                    for r in ws.iter_rows(values_only=True)]
            sheets.append((ws.title, [r for r in rows if any(r)]))
        return sheets
    text = read_csv_text(path)
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text),
                                                        delimiter=delimiter)]
    return [("", [r for r in rows if any(r)])]


def read_table(path):
    """Первый лист как список словарей {заголовок: значение}."""
    sheets = read_grid(path)
    if not sheets or not sheets[0][1]:
        return []
    head, *rows = sheets[0][1]
    return [dict(zip(head, r + [""] * (len(head) - len(r)))) for r in rows]


# Как называют колонки в чужих таблицах (списки с мероприятий и т.п.).
GUESS_HEADERS = {  # порядок важен: «Телеграм» раньше, чем «тел»
    "date": ("отметка времени", "дата", "время", "date", "time", "timestamp"),
    "telegram": ("telegram", "телеграм", "телега", "тг", "tg", "ник", "username",
                 "логин"),
    "phone": ("телефон", "phone", "номер", "тел", "mobile", "моб", "whatsapp",
              "ватсап", "сотов"),
    "name": ("фио", "имя", "name", "фамилия", "отчество", "контакт", "участник",
             "клиент", "гость"),
    "request": ("запрос", "интерес", "что ищете", "цель", "вопрос"),
    "sphere": ("сфера", "деятельн", "компания", "должность", "ниша", "бизнес",
               "профессия", "занима", "отрасль", "работ", "company", "position"),
}


def _guess_kind(header):
    h = header.lower()
    for kind, words in GUESS_HEADERS.items():
        if any(w in h for w in words):
            return kind
    return ""


def parse_any_sheet(rows):
    """Таблица неизвестного вида: ищем строку заголовков и колонки
    с телефоном / Telegram / именем / сферой; остальное — в заметки."""
    if not rows:
        return []
    head_i = next((i for i, r in enumerate(rows[:10])
                   if sum(bool(_guess_kind(c)) for c in r if c) >= 1
                   and not any(normalize_phone(c) for c in r)), None)
    width = max(len(r) for r in rows)
    body = rows[head_i + 1:] if head_i is not None else rows
    head = (rows[head_i] if head_i is not None else []) + [""] * width
    kinds = [_guess_kind(h) for h in head[:width]]
    # Колонка без понятного заголовка, но где в основном номера или @ники.
    for c in range(width):
        vals = [r[c] for r in body if c < len(r) and r[c]]
        if not vals or kinds[c]:
            continue
        if sum(bool(normalize_phone(v)) and 10 <= len(re.sub(r"\D", "", v)) <= 12
               and not re.search(r"\d{4}-\d\d-\d\d", v)
               for v in vals) > len(vals) / 2:
            kinds[c] = "phone"
        elif sum(v.startswith("@") or "t.me/" in v for v in vals) > len(vals) / 2:
            kinds[c] = "telegram"
    if "phone" not in kinds and "telegram" not in kinds:
        return []
    items = []
    for r in body:
        r = r + [""] * (width - len(r))
        item = {"phones": [], "name": "", "telegram": "", "sphere": "", "notes": "",
                "tags": [], "add_only": True}
        names, notes, spheres = [], [], []
        for c, kind in enumerate(kinds):
            v = r[c]
            if not v:
                continue
            if kind == "phone":
                item["phones"] += [p for p in re.split(r"[,;/]| {2,}", v) if p.strip()]
            elif kind == "telegram" and not item["telegram"]:
                if normalize_phone(v) and not v.startswith("@"):
                    item["phones"].append(v)  # в колонке «ТГ» бывает номер
                else:
                    nick = normalize_telegram(v.split()[0])
                    # Ник Telegram: латиница, цифры, _, от 5 символов.
                    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{4,31}", nick):
                        item["telegram"] = nick
            elif kind == "name":
                names.append(v)
            elif kind == "sphere":
                spheres.append(v)
            elif kind == "date":
                continue
            elif kind == "request":
                notes.append(f"Запрос: {v}")
                if not re.fullmatch(r"[-–—.\s]*(нет|пока нет|no|не знаю|-)?[\s).!]*",
                                    v.lower()):
                    item["tags"].append("есть-запрос")
            else:
                notes.append(f"{head[c]}: {v}" if head[c] else v)
        item["name"] = " ".join(names)
        item["sphere"] = ", ".join(spheres)
        item["notes"] = "; ".join(notes)
        if item["phones"] or item["telegram"]:
            items.append(item)
    return items


def parse_table(path):
    """Наша таблица (после export / списки чатов), экспорт Google Контактов
    или любая другая таблица с телефонами или ником Telegram."""
    table = read_table(path)
    head = {k.strip().lower() for k in (table[0] if table else {})}
    if not ({"phone", "телефон"} & head) and \
            not any(h.startswith("phone 1") for h in head):
        items = []
        for _, rows in read_grid(path):
            items += parse_any_sheet(rows)
        return items
    result = []
    for r in table:
        r = {k.strip(): (v or "").strip() for k, v in r.items()
             if isinstance(k, str) and isinstance(v, (str, type(None)))}
        mine = {HEADER_TO_KEY[k.lower()]: v for k, v in r.items()
                if k.lower() in HEADER_TO_KEY}
        if "phone" in mine:
            phone = mine.get("phone", "")
            if re.fullmatch(r"\d{10,15}(\.0)?", phone):  # Excel съел плюс
                phone = "+" + phone.split(".")[0]
            has_id = bool(re.sub(r"\.0$", "", mine.get("id", "")))
            result.append({
                "id": re.sub(r"\.0$", "", mine.get("id", "")),
                "phones": [phone],
                "name": mine.get("name", ""),
                "call_name": mine.get("call_name", ""),
                "gender": classify.parse_gender(mine.get("gender")),
                "region": mine.get("region", ""),
                "operator": mine.get("operator", ""),
                "sphere": mine.get("sphere", ""),
                "telegram": mine.get("telegram", ""),
                "tags": split_tags(mine.get("categories", "")) +
                        [chat_tag(t) for t in split_tags(mine.get("chats", ""))],
                "consent": yes_no(mine.get("consent")),
                "stopped": yes_no(mine.get("stopped")),
                "notes": mine.get("notes", ""),
                # С id — файл из export: он главный. Без id — список со стороны
                # (например, участники чата): только дополняет базу.
                "replace_tags": has_id,
                "add_only": not has_id,
            })
        else:  # Google Контакты
            name = r.get("Name") or " ".join(
                x for x in (r.get("First Name", ""), r.get("Last Name", "")) if x)
            phones = []
            for k, v in r.items():
                if k.startswith("Phone") and k.endswith("Value") and v:
                    phones += [p.strip() for p in v.split(":::")]
            labels = [t for t in split_tags(r.get("Labels", "").replace(":::", ","))
                      if t not in ("* mycontacts", "* starred")]
            result.append({"name": name, "phones": phones, "tags": labels,
                           "notes": r.get("Notes", ""), "in_phonebook": 1})
    return result


parse_csv = parse_table  # прежнее имя


TEXT_FIELDS = ("name", "call_name", "telegram", "notes", "gender", "region",
               "operator", "sphere")
INT_FIELDS = ("consent", "stopped", "in_phonebook")


def upsert_contact(db, item, extra_tags=()):
    """Добавляет контакт или обновляет существующий. Возвращает 'new'/'upd'/'skip'."""
    phones = [p for p in (normalize_phone(x) for x in item.get("phones", [])) if p]
    phone = phones[0] if phones else None
    row = None
    if item.get("id"):
        row = db.execute("SELECT * FROM contacts WHERE id=?", (item["id"],)).fetchone()
    if row is None and phone:
        row = db.execute("SELECT * FROM contacts WHERE phone=?", (phone,)).fetchone()
    tg = normalize_telegram(item.get("telegram"))
    if row is None and tg:
        row = db.execute("SELECT * FROM contacts WHERE telegram<>'' AND"
                         " lower(telegram)=lower(?)", (tg,)).fetchone()
    if row is None and not phone and not tg:
        return "skip"

    values = dict(item)
    values["telegram"] = normalize_telegram(item.get("telegram"))
    notes = item.get("notes", "")
    if len(phones) > 1:
        extra = "доп. номера: " + ", ".join(phones[1:])
        notes = f"{notes}; {extra}" if notes else extra
    values["notes"] = notes

    if row is None:
        fields = TEXT_FIELDS + INT_FIELDS
        db.execute(
            f"INSERT INTO contacts (phone, tags, {', '.join(fields)}, created_at,"
            f" updated_at) VALUES ({','.join('?' * (len(fields) + 4))})",
            (phone, join_tags(list(item.get("tags", [])) + list(extra_tags)),
             *[values.get(f) or "" for f in TEXT_FIELDS],
             *[int(values.get(f) or 0) for f in INT_FIELDS], now(), now()))
        item["_contact_id"] = db.execute("SELECT last_insert_rowid()").fetchone()[0]
        return "new"

    if item.get("replace_tags"):  # файл из export — это «истина» для тегов
        tags = list(item.get("tags", []))
    else:
        tags = split_tags(row["tags"]) + list(item.get("tags", []))
    tags += list(extra_tags)

    add_only = item.get("add_only")
    if add_only and notes and notes not in (row["notes"] or ""):
        values["notes"] = f"{row['notes']}; {notes}" if row["notes"] else notes
    elif add_only:
        values["notes"] = ""

    new = {}
    for f in TEXT_FIELDS + INT_FIELDS:
        value, old = values.get(f), row[f]
        if f == "notes" and add_only and value:
            new[f] = value  # склеенная заметка
        elif f == "call_name" and add_only and row["name"]:
            new[f] = old  # обращение — по вашей записи, не по нику из чата
        elif value in (None, "") or (add_only and old not in (None, "", 0)):
            new[f] = old  # не затираем имя из телефонной книги и т.п.
        else:
            new[f] = value
    new["in_phonebook"] = max(int(row["in_phonebook"] or 0),
                              int(values.get("in_phonebook") or 0))
    fields = list(new)
    db.execute(
        f"UPDATE contacts SET phone=?, tags=?, {', '.join(f + '=?' for f in fields)},"
        " updated_at=? WHERE id=?",
        (phone or row["phone"], join_tags(tags),
         *[int(new[f] or 0) if f in INT_FIELDS else new[f] for f in fields],
         now(), row["id"]))
    item["_contact_id"] = row["id"]
    return "upd"


# ---------------------------------------------------------------- выборка, шаблоны

def find_contact(db, key):
    if str(key).startswith("@"):
        row = db.execute("SELECT * FROM contacts WHERE lower(telegram)=lower(?)",
                         (normalize_telegram(key),)).fetchone()
    elif str(key).isdigit() and len(str(key)) < 7:
        row = db.execute("SELECT * FROM contacts WHERE id=?", (int(key),)).fetchone()
    else:
        row = db.execute("SELECT * FROM contacts WHERE phone=?",
                         (normalize_phone(key),)).fetchone()
    if row is None:
        sys.exit(f"Контакт «{key}» не найден")
    return row


def select_contacts(db, tags=(), exclude=(), search="", gender="", region=""):
    rows = db.execute("SELECT * FROM contacts ORDER BY name='', name").fetchall()
    gender = classify.parse_gender(gender) if gender else ""
    out = []
    for r in rows:
        # Сфера деятельности тоже работает как категория в --tags.
        ctags = set(split_tags(r["tags"])) | set(split_tags(r["sphere"]))
        if gender and r["gender"] != gender:
            continue
        if region and region.lower() not in (r["region"] or "").lower():
            continue
        if tags and not ctags & set(tags):
            continue
        if exclude and ctags & set(exclude):
            continue
        if search and search.lower() not in (
                f"{r['name']} {r['phone']} {r['notes']} {r['telegram']}".lower()):
            continue
        out.append(r)
    return out


class SafeDict(dict):
    def __missing__(self, key):
        return ""


def greeting(moment=None):
    h = (moment or dt.datetime.now()).hour
    if 5 <= h < 12:
        return "Доброе утро"
    if 12 <= h < 18:
        return "Добрый день"
    return "Добрый вечер"


def render(template, contact):
    name = (contact["name"] or "").strip()
    keys = contact.keys() if hasattr(contact, "keys") else ()
    first = (contact["call_name"] if "call_name" in keys else "") \
        or guess_first_name(name)
    return template.format_map(SafeDict(
        name=name, first_name=first, hello=greeting(),
        phone=contact["phone"] or "", notes=contact["notes"] or "",
    )).replace(", !", "!").strip()


# ---------------------------------------------------------------- каналы

class StopChannel(Exception):
    """Канал нужно остановить на этот запуск (лимит, бан, нет настроек)."""


class TelegramSender:
    """Отправка от вашего личного аккаунта Telegram (Telethon)."""

    def __init__(self, cfg):
        try:
            from telethon.sync import TelegramClient
        except ImportError:
            raise StopChannel("не установлен telethon: pip install -r requirements.txt")
        if not telegram_configured(cfg):
            raise StopChannel("Telegram не настроен: заполните [telegram] в config.ini "
                              "и выполните python3 crm.py tg-login")
        t = cfg["telegram"]
        self.client = TelegramClient(str(DATA_DIR / "telegram"),
                                     int(t["api_id"]), t["api_hash"])
        # При первом запуске Telegram пришлёт код входа — введите его в консоли.
        self.client.start(phone=t.get("phone"))

    def _resolve(self, contact):
        from telethon.tl.functions.contacts import ImportContactsRequest
        from telethon.tl.types import InputPhoneContact
        if contact["telegram"]:
            return self.client.get_entity(contact["telegram"])
        res = self.client(ImportContactsRequest([InputPhoneContact(
            client_id=random.randrange(1 << 62), phone=contact["phone"],
            first_name=contact["name"] or contact["phone"], last_name="")]))
        return res.users[0] if res.users else None

    def send(self, contact, text, photo=""):
        from telethon import errors
        try:
            user = self._resolve(contact)
            if user is None:
                return "no_account", "номер не зарегистрирован в Telegram"
            if photo and len(text) <= 1024:
                self.client.send_file(user, photo, caption=text)
            else:
                if photo:
                    self.client.send_file(user, photo)
                self.client.send_message(user, text)
            return "sent", ""
        except errors.FloodWaitError as e:
            raise StopChannel(f"Telegram просит подождать {e.seconds} с — остановлено")
        except errors.PeerFloodError:
            raise StopChannel("Telegram ограничил аккаунт за рассылку (PeerFlood) — "
                              "остановлено, сделайте паузу на сутки")
        except (ValueError, errors.RPCError) as e:
            return "error", str(e)

    def history(self, contact, limit=30):
        """Последние сообщения переписки с человеком: [(id, дата, out?, текст)]."""
        user = self._resolve(contact)
        if user is None:
            return []
        return [(m.id, m.date.astimezone().replace(tzinfo=None).isoformat(timespec="seconds"),
                 m.out, m.raw_text or "[вложение]")
                for m in self.client.iter_messages(user, limit=limit)]

    def stopped_senders(self, days=30):
        """Номера/юзернеймы тех, кто ответил «стоп» за последние дни."""
        since = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=days)
        found = []
        for d in self.client.iter_dialogs():
            if not d.is_user or d.date < since:
                continue
            for m in self.client.iter_messages(d.entity, limit=20):
                if m.date < since:
                    break
                if not m.out and re.search(r"\b(стоп|stop|отпишите)\b",
                                           m.raw_text or "", re.I):
                    found.append((getattr(d.entity, "phone", None),
                                  getattr(d.entity, "username", None)))
                    break
        return found


class SmsSender:
    """SMS с вашей SIM-карты через приложение «SMS Gateway for Android»."""

    def __init__(self, cfg):
        import requests
        if not cfg.has_section("sms"):
            raise StopChannel("нет секции [sms] в config.ini")
        s = cfg["sms"]
        self.url = s["url"].rstrip("/") + "/message"
        self.auth = (s["username"], s["password"])
        self.requests = requests

    def send(self, contact, text, photo=""):
        if photo:
            text = f"{text}\n{photo}"
        try:
            r = self.requests.post(self.url, auth=self.auth, timeout=30, json={
                "message": text, "phoneNumbers": [contact["phone"]]})
        except self.requests.RequestException as e:
            raise StopChannel(f"телефон-шлюз недоступен: {e}")
        if r.status_code in (401, 403):
            raise StopChannel("шлюз SMS: неверный логин/пароль")
        if r.status_code >= 400:
            return "error", f"HTTP {r.status_code}: {r.text[:200]}"
        return "queued", ""


class MacMessagesSender:
    """SMS с вашего номера через «Сообщения» на Mac (iPhone пересылает SMS:
    Настройки iPhone → Сообщения → Переадресация → включить этот Mac)."""

    SCRIPTS = (
        # macOS 11+
        '''on run argv
  tell application "Messages"
    set s to 1st account whose service type = SMS
    send (item 2 of argv) to participant (item 1 of argv) of s
  end tell
end run''',
        # старые macOS
        '''on run argv
  tell application "Messages"
    set s to 1st service whose service type = SMS
    send (item 2 of argv) to buddy (item 1 of argv) of s
  end tell
end run''',
    )

    def __init__(self, cfg=None):
        if sys.platform != "darwin":
            raise StopChannel("SMS через «Сообщения» работает только на Mac")

    def send(self, contact, text, photo=""):
        import subprocess
        if photo:
            text = f"{text}\n{photo}"
        err = ""
        for script in self.SCRIPTS:
            r = subprocess.run(["osascript", "-e", script, contact["phone"], text],
                               capture_output=True, text=True, timeout=60)
            if r.returncode == 0:
                return "sent", ""
            err = r.stderr.strip()
        return "error", err[:300]


class WhatsAppLinks:
    """WhatsApp: личный номер нельзя безопасно автоматизировать (бан за ботов),
    поэтому готовим страницу со ссылками wa.me — текст уже вставлен,
    остаётся нажать «Отправить» в WhatsApp у каждого контакта."""

    def __init__(self, campaign_name):
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
        safe = re.sub(r"[^\w-]+", "_", campaign_name)
        self.path = DATA_DIR / f"wa_links_{safe}_{stamp}.html"
        self.title = campaign_name
        self.items = []

    def send(self, contact, text, photo=""):
        if photo:
            text = f"{text}\n{photo}"
        digits = re.sub(r"\D", "", contact["phone"])
        url = f"https://wa.me/{digits}?text={urllib.parse.quote(text)}"
        self.items.append((contact["name"] or contact["phone"], contact["phone"], url))
        return "link", ""

    def close(self):
        if not self.items:
            return None
        rows = "\n".join(
            f'<li><label><input type="checkbox"> <a href="{html.escape(u)}" '
            f'target="_blank">{html.escape(n)}</a> <small>{html.escape(p)}</small>'
            f"</label></li>" for n, p, u in self.items)
        self.path.write_text(
            "<!doctype html><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width'>"
            f"<title>WhatsApp: {html.escape(self.title)}</title>"
            "<style>body{font:18px sans-serif;margin:16px}li{margin:14px 0}</style>"
            f"<h2>{html.escape(self.title)} — {len(self.items)} чел.</h2>"
            "<p>Откройте на телефоне, нажимайте имя → «Отправить» в WhatsApp, "
            "отмечайте галочкой.</p>"
            f"<ol>{rows}</ol>", encoding="utf-8")
        return self.path


# ---------------------------------------------------------------- команды

def cmd_import(db, a):
    path = Path(a.file)
    if not path.exists():
        sys.exit(f"Файл не найден: {path}")
    if path.suffix.lower() == ".vcf":
        items = parse_vcf(path.read_text(encoding="utf-8", errors="replace"))
        for it in items:
            it["in_phonebook"] = 1
    elif path.suffix.lower() in (".csv", ".xlsx"):
        items = parse_table(path)
    else:
        sys.exit("Поддерживаются .vcf (из телефона), .csv и .xlsx")
    stats = {"new": 0, "upd": 0, "skip": 0}
    source = path.name
    for item in items:
        status = upsert_contact(db, item, split_tags(a.tag))
        stats[status] += 1
        # В историю: откуда человек появился или что о нём добавилось.
        if status == "new":
            log_interaction(db, item["_contact_id"], "добавлен в базу",
                            f"из файла {source}", ext_id=f"import:{source}")
        elif status == "upd" and item.get("add_only"):
            log_interaction(db, item["_contact_id"], "найден в списке",
                            f"{source}", ext_id=f"import:{source}")
    db.commit()
    print(f"Новых: {stats['new']}, обновлено: {stats['upd']}, "
          f"без номера (пропущено): {stats['skip']}")


def contact_view(r):
    """Строка таблицы для человека: теги делятся на категории и чаты."""
    tags = split_tags(r["tags"])
    chats = [chat_title(t) for t in tags if is_source(t)]
    sources = (["телефон"] if r["in_phonebook"] else []) + chats
    return {
        "id": r["id"], "phone": r["phone"], "name": r["name"],
        "call_name": r["call_name"] or guess_first_name(r["name"]),
        "gender": classify.GENDER_LABEL.get(r["gender"], ""),
        "region": r["region"], "operator": r["operator"], "sphere": r["sphere"],
        "categories": ", ".join(t for t in tags if not is_source(t)),
        "chats": ", ".join(chats),
        "sources": " + ".join(sources),
        # Пересечение — человек есть и у вас в телефоне, и в чате/списке.
        "overlap": "да" if r["in_phonebook"] and chats else "",
        "consent": r["consent"], "stopped": r["stopped"],
        "telegram": r["telegram"], "notes": r["notes"],
    }


def cmd_export(db, a):
    rows = [contact_view(r) for r in select_contacts(
        db, split_tags(a.tag), gender=a.gender, region=a.region)]
    if a.overlap:
        rows = [r for r in rows if r["overlap"]]
    head = [h for _, h in COLUMNS]
    keys = [k for k, _ in COLUMNS]
    if a.file.lower().endswith(".xlsx"):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            sys.exit("Для .xlsx установите: python3 -m pip install --user openpyxl")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Контакты"
        ws.append(head)
        for r in rows:
            ws.append([r[k] for k in keys])
        widths = {"id": 6, "phone": 15, "name": 32, "call_name": 14, "gender": 10,
                  "region": 28, "operator": 20, "sphere": 22, "categories": 22,
                  "chats": 22, "sources": 30, "overlap": 12, "consent": 10,
                  "stopped": 7, "telegram": 16, "notes": 50}
        for i, k in enumerate(keys, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths[k]
            ws.cell(1, i).font = Font(bold=True)
        fill = PatternFill("solid", fgColor="FFF2CC")
        ov = keys.index("overlap") + 1
        for row in ws.iter_rows(min_row=2):
            if row[ov - 1].value:
                for c in row:
                    c.fill = fill  # пересечения подсвечены
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(a.file)
    else:
        with open(a.file, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(head)
            for r in rows:
                w.writerow([r[k] for k in keys])
    print(f"Выгружено {len(rows)} контактов в {a.file}. Можно править: Сфера "
          "деятельности, Категории, Пол, Обращение, Согласие (1 — согласен). "
          f"Потом загрузить обратно: python3 crm.py import {a.file}")


def cmd_list(db, a):
    rows = select_contacts(db, split_tags(a.tag), search=a.search or "",
                           gender=a.gender, region=a.region)
    for r in rows:
        flags = ("" if r["consent"] else " [нет согласия]") + \
                (" [СТОП]" if r["stopped"] else "") + \
                ("" if phone_ok(r["phone"]) or (r["telegram"] and not r["phone"])
                 else " [ОШИБКА В НОМЕРЕ]")
        tg = f" @{r['telegram']}" if r["telegram"] else ""
        print(f"{r['id']:>5}  {r['phone'] or '':<14} {r['name'][:28]:<28} "
              f"{r['gender'] or '?'}  {r['region'][:18]:<18} "
              f"{r['tags']}{tg}{flags}")
    print(f"Итого: {len(rows)}")


def cmd_tags(db, a):
    counts = {}
    for r in db.execute("SELECT tags FROM contacts"):
        for t in split_tags(r["tags"]):
            counts[t] = counts.get(t, 0) + 1
    for t, c in sorted(counts.items(), key=lambda x: -x[1]):
        print(f"{c:>6}  {t}")


def cmd_words(db, a):
    """Частые слова в именах — так видно, какими пометками вы подписывали людей."""
    counts = {}
    for r in db.execute("SELECT name FROM contacts"):
        for w in set(re.findall(r"[A-Za-zА-Яа-яЁё]{3,}", r["name"] or "")):
            w = w.lower()
            counts[w] = counts.get(w, 0) + 1
    top = sorted(counts.items(), key=lambda x: -x[1])[: a.top]
    for w, c in top:
        print(f"{c:>6}  {w}")


def cmd_autotag(db, a):
    """Всем, у кого в имени или заметке есть слово, добавить тег."""
    words = [w.lower() for w in a.words.split(",") if w.strip()]
    tag = a.tag.strip().lower()
    n = 0
    for r in db.execute("SELECT * FROM contacts").fetchall():
        text = f"{r['name']} {r['notes']}".lower()
        if not any(w.strip() in text for w in words):
            continue
        tags = split_tags(r["tags"])
        if tag in tags:
            continue
        n += 1
        if a.dry:
            print(f"  {r['phone']:<14} {r['name']}")
            continue
        db.execute("UPDATE contacts SET tags=?, updated_at=? WHERE id=?",
                   (join_tags(tags + [tag]), now(), r["id"]))
    db.commit()
    print(("Будет отмечено" if a.dry else "Отмечено") + f" тегом «{tag}»: {n}")


def cmd_regions(db, a):
    folder = DATA_DIR / "numbering"
    if a.dir:
        import shutil
        folder.mkdir(parents=True, exist_ok=True)
        n = 0
        for name in classify.REGISTRY_FILES:
            src = Path(a.dir).expanduser() / name
            if src.exists():
                shutil.copy(src, folder / name)
                n += 1
        print(f"Скопировано файлов реестра: {n}")
    else:
        print("Скачиваю реестр номеров Минцифры (несколько минут)…")
        n = classify.download_registry(folder, insecure=a.insecure)
    if not classify.Numbering(folder):
        sys.exit("Реестр не загружен. Скачайте в браузере файлы DEF-9xx.csv "
                 f"(и ABC-3xx/4xx/8xx) со страницы {classify.REGISTRY_PAGE} "
                 "и выполните: python3 crm.py regions --dir ~/Downloads")
    print("Реестр готов. Теперь: python3 crm.py classify")


def cmd_classify(db, a):
    """Регион, пол, категории по пометкам — для всей базы."""
    numbering = classify.Numbering(DATA_DIR / "numbering")
    rules = classify.load_rules(RULES_PATH)
    stats = {"region": 0, "gender": 0, "tags": 0}
    for r in db.execute("SELECT * FROM contacts").fetchall():
        region, operator = classify.region_of(r["phone"], numbering)
        gender = r["gender"] or classify.guess_gender(r["call_name"], r["name"])
        tags = split_tags(r["tags"])
        added = [t for t in classify.match_rules(rules, r["name"], r["notes"], tags)
                 if t not in tags]
        if region and (region, operator) != (r["region"], r["operator"]):
            stats["region"] += 1
        else:
            region, operator = r["region"], r["operator"]
        stats["gender"] += gender != r["gender"]
        stats["tags"] += bool(added)
        db.execute("UPDATE contacts SET region=?, operator=?, gender=?, tags=?,"
                   " updated_at=? WHERE id=?",
                   (region, operator, gender, join_tags(tags + added), now(), r["id"]))
    db.commit()

    def count(sql, *args):
        return db.execute(sql, args).fetchone()[0]

    total = count("SELECT COUNT(*) FROM contacts")
    print(f"Обновлено: регион у {stats['region']}, пол у {stats['gender']}, "
          f"категории у {stats['tags']}.")
    print(f"Всего {total}: женщин {count('SELECT COUNT(*) FROM contacts WHERE gender=?', 'ж')}, "
          f"мужчин {count('SELECT COUNT(*) FROM contacts WHERE gender=?', 'м')}, "
          f"пол не определён {count('SELECT COUNT(*) FROM contacts WHERE gender=?', '')}.")
    if not numbering:
        print("Регион по России пока общий («Россия»). Для точного региона "
              "и оператора выполните: python3 crm.py regions")
    print()
    cmd_tags(db, a)
    overlap = sum(1 for r in db.execute("SELECT * FROM contacts")
                  if contact_view(r)["overlap"])
    print(f"\nПересечений (есть и у вас в телефоне, и в чатах/списках): {overlap}")


def cmd_set(db, a):
    r = find_contact(db, a.contact)
    tags = [t for t in split_tags(r["tags"]) if t not in split_tags(a.remove_tags)]
    tags += split_tags(a.add_tags)
    notes = r["notes"]
    if a.note:
        notes = f"{notes}\n{dt.date.today()}: {a.note}".strip()
        log_interaction(db, r["id"], "заметка", a.note)
    db.execute(
        "UPDATE contacts SET tags=?, telegram=?, consent=?, stopped=?, notes=?,"
        " updated_at=? WHERE id=?",
        (join_tags(tags),
         r["telegram"] if a.telegram is None else normalize_telegram(a.telegram),
         r["consent"] if a.consent is None else a.consent,
         r["stopped"] if a.stop is None else a.stop,
         notes, now(), r["id"]))
    db.commit()
    cmd_history(db, argparse.Namespace(contact=a.contact))


def cmd_history(db, a):
    r = find_contact(db, a.contact)
    print(f"{r['name']}  {r['phone'] or ('@' + r['telegram'])}  теги: {r['tags'] or '—'}  "
          f"согласие: {'да' if r['consent'] else 'нет'}"
          f"{'  [СТОП]' if r['stopped'] else ''}")
    if r["notes"]:
        print(f"Заметки: {r['notes']}")
    for e in timeline(db, r["id"]):
        arrow = {"out": "→", "in": "←"}.get(e["direction"], "•")
        print(f"  {e['at'][:16].replace('T', ' ')}  {arrow} "
              f"{CHANNEL_TITLES.get(e['channel'], e['channel']):<9} {e['kind']}"
              f"{': ' + e['text'][:80] if e['text'] else ''}"
              f"{' [' + e['status'] + ']' if e['status'] else ''}")


def cmd_campaigns(db, a):
    for c in db.execute("SELECT * FROM campaigns ORDER BY id"):
        stats = dict(db.execute(
            "SELECT status, COUNT(*) FROM messages WHERE campaign_id=?"
            " GROUP BY status", (c["id"],)).fetchall())
        print(f"{c['id']:>3}  {c['created_at'][:16]}  {c['name']:<30} "
              f"теги: {c['tags'] or 'все'}  {stats}")


def cmd_report(db, a):
    c = db.execute("SELECT * FROM campaigns WHERE name=? OR id=?",
                   (a.campaign, a.campaign)).fetchone()
    if c is None:
        sys.exit("Кампания не найдена")
    rows = db.execute(
        "SELECT m.*, k.name, k.phone FROM messages m JOIN contacts k"
        " ON k.id=m.contact_id WHERE campaign_id=? ORDER BY m.sent_at", (c["id"],))
    lines = [[m["name"], m["phone"], m["channel"], m["status"], m["error"],
              m["sent_at"]] for m in rows]
    if not a.file:
        for line in lines:
            print("  ".join(str(x) for x in line))
        return
    with open(a.file, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["name", "phone", "channel", "status", "error", "sent_at"])
        w.writerows(lines)
    print(f"Сохранено: {a.file}")


def cmd_sync_stop(db, a):
    cfg = load_config()
    try:
        tg = TelegramSender(cfg)
    except StopChannel as e:
        sys.exit(f"Telegram: {e}")
    n = 0
    for phone, username in tg.stopped_senders(a.days):
        cur = db.execute(
            "UPDATE contacts SET stopped=1, updated_at=? WHERE stopped=0 AND"
            " (phone=? OR (telegram<>'' AND lower(telegram)=lower(?)))",
            (now(), normalize_phone(phone and "+" + phone), username or ""))
        n += cur.rowcount
    db.commit()
    print(f"Отмечено «СТОП»: {n}")


def cmd_send(db, a):
    template = Path(a.template).read_text(encoding="utf-8")
    channels = [c.strip() for c in a.channels.split(",") if c.strip()]
    bad = [c for c in channels if c not in CHANNELS]
    if bad:
        sys.exit(f"Неизвестный канал: {bad}. Доступны: {', '.join(CHANNELS)}")
    tags, exclude = split_tags(a.tags), split_tags(a.exclude_tags)

    everyone = select_contacts(db, tags, exclude, gender=a.gender, region=a.region)
    stopped = [r for r in everyone if r["stopped"]]
    no_consent = [r for r in everyone if not r["stopped"] and not r["consent"]]
    bad_phone = [r for r in everyone
                 if not phone_ok(r["phone"]) and not r["telegram"]]
    targets = [r for r in everyone
               if not r["stopped"] and r["consent"]
               and (phone_ok(r["phone"]) or r["telegram"])]

    campaign = db.execute("SELECT * FROM campaigns WHERE name=?", (a.name,)).fetchone()
    if campaign:
        done = {row[0] for row in db.execute(
            "SELECT contact_id FROM messages WHERE campaign_id=? AND status IN"
            f" ({','.join('?' * len(DELIVERED))})", (campaign["id"], *DELIVERED))}
        targets = [r for r in targets if r["id"] not in done]

    print(f"Кампания «{a.name}»: подходит по тегам {len(everyone)}, "
          f"отписались {len(stopped)}, без согласия {len(no_consent)}, "
          f"с ошибкой в номере {len(bad_phone)} (все они пропущены), "
          f"к отправке {len(targets)}. Каналы по порядку: {' → '.join(channels)}")
    if targets:
        print("\n--- пример сообщения для", targets[0]["name"] or targets[0]["phone"])
        print(render(template, targets[0]))
        print("---")

    if not a.send:
        for r in targets[: a.limit or len(targets)]:
            print(f"  {r['phone']:<14} {r['name']}")
        print("\nЭто пробный прогон, ничего не отправлено. "
              "Для отправки добавьте --send")
        return

    if campaign is None:
        db.execute(
            "INSERT INTO campaigns (name, template, photo, tags, channels, created_at)"
            " VALUES (?,?,?,?,?,?)",
            (a.name, template, a.photo or "", join_tags(tags), ",".join(channels), now()))
        db.commit()
        campaign = db.execute("SELECT * FROM campaigns WHERE name=?",
                              (a.name,)).fetchone()

    cfg = load_config()
    lim = cfg["limits"] if cfg.has_section("limits") else {}
    min_d, max_d = float(lim.get("min_delay", 20)), float(lim.get("max_delay", 60))
    per_run = a.limit or int(lim.get("per_run", 40))

    senders, dead, sent_count = {}, {}, {c: 0 for c in channels}

    def get_sender(ch):
        if ch in dead:
            return None
        if ch not in senders:
            try:
                senders[ch] = {"tg": lambda: TelegramSender(cfg),
                               "sms": lambda: SmsSender(cfg),
                               "wa": lambda: WhatsAppLinks(a.name)}[ch]()
            except StopChannel as e:
                dead[ch] = str(e)
                print(f"[{ch}] отключён: {e}")
                return None
        return senders[ch]

    def log(contact, ch, status, err, text):
        db.execute(
            "INSERT INTO messages (campaign_id, contact_id, channel, status, error,"
            " text, sent_at) VALUES (?,?,?,?,?,?,?) ON CONFLICT(campaign_id,"
            " contact_id, channel) DO UPDATE SET status=excluded.status,"
            " error=excluded.error, text=excluded.text, sent_at=excluded.sent_at",
            (campaign["id"], contact["id"], ch, status, err, text, now()))
        db.commit()

    try:
        for r in targets:
            text = render(template, r)
            for ch in channels:
                if ch in dead or sent_count[ch] >= per_run:
                    continue
                if ch != "tg" and not phone_ok(r["phone"]):
                    continue  # только ник Telegram, номера нет
                sender = get_sender(ch)
                if sender is None:
                    continue
                try:
                    status, err = sender.send(r, text, a.photo or "")
                except StopChannel as e:
                    dead[ch] = str(e)
                    print(f"[{ch}] остановлен: {e}")
                    continue
                log(r, ch, status, err, text)
                print(f"[{ch}] {r['phone']:<14} {r['name'][:30]:<30} {status} {err}")
                if status in DELIVERED:
                    sent_count[ch] += 1
                    if ch != "wa":
                        time.sleep(random.uniform(min_d, max_d))
                    break  # один человек — одно сообщение (первый сработавший канал)
            if all(ch in dead or sent_count[ch] >= per_run for ch in channels):
                print("Лимиты на этот запуск исчерпаны — запустите позже, "
                      "уже отправленным повторно не придёт.")
                break
    except KeyboardInterrupt:
        print("\nПрервано. Отправленное сохранено в базе.")
    finally:
        wa = senders.get("wa")
        page = wa.close() if wa else None
        if page:
            print(f"\nWhatsApp: откройте {page} и отправьте по ссылкам.")
        tg = senders.get("tg")
        if tg:
            tg.client.disconnect()
    print("Итог по каналам:", sent_count)


def cmd_tg_login(db, a):
    """Первый вход в Telegram: код из приложения Telegram (и облачный пароль,
    если он включён) вводится здесь, дальше вход запоминается."""
    cfg = load_config()
    if not telegram_configured(cfg) or a.reset:
        print("Нужны ключи с https://my.telegram.org → API development tools "
              "(App api_id и App api_hash).")
        api_id = input("api_id (только цифры): ").strip()
        api_hash = input("api_hash (32 символа): ").strip()
        phone = normalize_phone(input("Ваш номер телефона в Telegram: ").strip())
        if not cfg.has_section("telegram"):
            cfg.add_section("telegram")
        cfg["telegram"].update(api_id=api_id, api_hash=api_hash, phone=phone or "")
        if not telegram_configured(cfg):
            sys.exit("Ключи не похожи на настоящие: api_id — только цифры, "
                     "api_hash — 32 символа. Запустите команду ещё раз.")
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            cfg.write(f)
        print("Ключи сохранены в config.ini (он только на вашем компьютере).")
    print("Подключаюсь к Telegram. Если попросит — введите код, который придёт "
          "в приложение Telegram (не SMS), и облачный пароль, если он у вас есть.")
    try:
        tg = TelegramSender(cfg)
    except StopChannel as e:
        sys.exit(str(e))
    me = tg.client.get_me()
    tg.client.send_message("me", "✅ База контактов подключена к Telegram. "
                                 "Отсюда можно отправлять сообщения.")
    print(f"Готово: вход как {me.first_name or ''} {me.last_name or ''} (+{me.phone}). "
          "Проверочное сообщение — в «Избранном» в Telegram.")
    tg.client.disconnect()


def cmd_shortcut(db, a):
    """Значок «База контактов» на рабочем столе: двойной щелчок — запуск app."""
    desktop = Path.home() / "Desktop"
    if not desktop.exists():
        desktop = Path.home()
    path = desktop / "База контактов.command"
    path.write_text(
        "#!/bin/bash\n"
        f'cd "{BASE_DIR}" || exit 1\n'
        "echo 'Открываю базу контактов… Не закрывайте это окно, пока работаете.'\n"
        "python3 crm.py app\n", encoding="utf-8")
    path.chmod(0o755)
    print(f"Готово: на рабочем столе значок «{path.stem}». "
          "Двойной щелчок — откроется база в браузере.")


def cmd_app(db, a):
    import webapp
    webapp.run(db, port=a.port, open_browser=not a.no_browser)


def build_parser():
    p = argparse.ArgumentParser(description="CRM брокера: контакты и рассылки")
    p.add_argument("--db", help="путь к базе (по умолчанию data/crm.db)")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("import", help="загрузить контакты из .vcf или .csv")
    s.add_argument("file")
    s.add_argument("--tag", default="", help="добавить всем этот тег(и)")
    s.set_defaults(func=cmd_import)

    s = sub.add_parser("export", help="выгрузить контакты в .xlsx или .csv")
    s.add_argument("file")
    s.add_argument("--tag", default="")
    s.add_argument("--gender", default="", help="ж или м")
    s.add_argument("--region", default="", help="часть названия региона")
    s.add_argument("--overlap", action="store_true",
                   help="только те, кто есть и в телефоне, и в чатах")
    s.set_defaults(func=cmd_export)

    s = sub.add_parser("list", help="показать контакты")
    s.add_argument("--tag", default="")
    s.add_argument("--gender", default="")
    s.add_argument("--region", default="")
    s.add_argument("--search", default="")
    s.set_defaults(func=cmd_list)

    sub.add_parser("tags", help="категории и число людей в них") \
        .set_defaults(func=cmd_tags)

    sub.add_parser("classify", help="регион, пол и категории для всей базы") \
        .set_defaults(func=cmd_classify)

    s = sub.add_parser("regions", help="загрузить реестр номеров (регион, оператор)")
    s.add_argument("--dir", help="папка, куда вы сами скачали DEF-9xx.csv")
    s.add_argument("--insecure", action="store_true",
                   help="не проверять сертификат сайта Минцифры (реестр — открытые данные)")
    s.set_defaults(func=cmd_regions)

    s = sub.add_parser("words", help="частые слова в именах контактов")
    s.add_argument("--top", type=int, default=80)
    s.set_defaults(func=cmd_words)

    s = sub.add_parser("autotag", help="тег всем, у кого в имени есть слово")
    s.add_argument("words", help="слово или несколько через запятую: покупат,купить")
    s.add_argument("tag", help="какой тег поставить")
    s.add_argument("--dry", action="store_true", help="только показать, кого отметит")
    s.set_defaults(func=cmd_autotag)

    s = sub.add_parser("set", help="изменить контакт (по номеру или id)")
    s.add_argument("contact")
    s.add_argument("--add-tags", default="")
    s.add_argument("--remove-tags", default="")
    s.add_argument("--telegram")
    s.add_argument("--consent", type=int, choices=(0, 1))
    s.add_argument("--stop", type=int, choices=(0, 1))
    s.add_argument("--note", help="добавить заметку с датой")
    s.set_defaults(func=cmd_set)

    s = sub.add_parser("history", help="карточка и история по контакту")
    s.add_argument("contact")
    s.set_defaults(func=cmd_history)

    s = sub.add_parser("send", help="рассылка (по умолчанию — пробный прогон)")
    s.add_argument("--name", required=True, help="название кампании")
    s.add_argument("--template", required=True, help="файл с текстом")
    s.add_argument("--tags", default="",
                   help="кому: категории или сферы через запятую (любая из них)")
    s.add_argument("--gender", default="", help="ж или м")
    s.add_argument("--region", default="", help="часть названия региона, напр. Москва")
    s.add_argument("--exclude-tags", default="")
    s.add_argument("--channels", default="tg,wa",
                   help="порядок каналов, напр. tg,sms,wa")
    s.add_argument("--photo", default="", help="ссылка на фото объекта")
    s.add_argument("--limit", type=int, default=0,
                   help="максимум сообщений на канал за запуск")
    s.add_argument("--send", action="store_true", help="реально отправить")
    s.set_defaults(func=cmd_send)

    sub.add_parser("shortcut", help="значок запуска на рабочем столе") \
        .set_defaults(func=cmd_shortcut)

    s = sub.add_parser("tg-login", help="подключить ваш Telegram (один раз)")
    s.add_argument("--reset", action="store_true", help="ввести ключи заново")
    s.set_defaults(func=cmd_tg_login)

    s = sub.add_parser("app", help="открыть базу в браузере: карточки, история, сообщения")
    s.add_argument("--port", type=int, default=8765)
    s.add_argument("--no-browser", action="store_true")
    s.set_defaults(func=cmd_app)

    sub.add_parser("campaigns", help="список рассылок и статистика") \
        .set_defaults(func=cmd_campaigns)

    s = sub.add_parser("report", help="кому и как ушла кампания")
    s.add_argument("campaign")
    s.add_argument("--file", help="сохранить в .csv")
    s.set_defaults(func=cmd_report)

    s = sub.add_parser("sync-stop", help="отметить ответивших «СТОП» в Telegram")
    s.add_argument("--days", type=int, default=30)
    s.set_defaults(func=cmd_sync_stop)
    return p


def main(argv=None):
    a = build_parser().parse_args(argv)
    db = connect(a.db)
    try:
        a.func(db, a)
    finally:
        db.close()


if __name__ == "__main__":
    main()
