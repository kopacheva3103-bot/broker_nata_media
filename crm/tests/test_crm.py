import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import crm  # noqa: E402

VCF = """BEGIN:VCARD
VERSION:2.1
N;CHARSET=UTF-8;ENCODING=QUOTED-PRINTABLE:=D0=9F=D0=B5=D1=82=D1=80=D0=BE=D0=B2;=D0=98=D0=B2=D0=B0=D0=BD;;;
TEL;CELL:8 (916) 123-45-67
TEL;HOME:+7 495 000-00-00
END:VCARD
BEGIN:VCARD
VERSION:3.0
FN:Ольга Смирнова
item1.TEL;type=CELL:+7 903 111-22-33
CATEGORIES:Инвестор,VIP
NOTE:Коммерция
END:VCARD
BEGIN:VCARD
VERSION:3.0
FN:Без номера
END:VCARD
"""


class CrmTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.db = crm.connect(self.dir / "t.db")

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def test_normalize_phone(self):
        self.assertEqual(crm.normalize_phone("8 (916) 123-45-67"), "+79161234567")
        self.assertEqual(crm.normalize_phone("9161234567"), "+79161234567")
        self.assertEqual(crm.normalize_phone("+375 29 123 45 67"), "+375291234567")
        self.assertIsNone(crm.normalize_phone("112"))
        self.assertIsNone(crm.normalize_phone("+7 (932) 300-57-8"))
        self.assertFalse(crm.phone_ok("+7932300578"))

    def test_guess_first_name(self):
        self.assertEqual(crm.guess_first_name(". Светлана Владимировна Мира 1"), "Светлана")
        self.assertEqual(crm.guess_first_name("Галина Вайбер Зал"), "Галина")
        self.assertEqual(crm.guess_first_name("2ой Покупатель Снт Урал"), "")
        self.assertEqual(crm.guess_first_name("+7 (912) 083-27-79"), "")
        self.assertEqual(crm.guess_first_name("Покупатель Иван"), "Иван")
        self.assertEqual(crm.guess_first_name("Покупатель Снт"), "")

    def test_autotag(self):
        for it in crm.parse_vcf(VCF):
            crm.upsert_contact(self.db, it)
        crm.cmd_autotag(self.db, type("A", (), {"words": "петров", "tag": "Покупатель",
                                                "dry": False}))
        r = self.db.execute("SELECT tags FROM contacts WHERE phone='+79161234567'").fetchone()
        self.assertEqual(r["tags"], "покупатель")

    def test_parse_vcf(self):
        items = crm.parse_vcf(VCF)
        self.assertEqual(items[0]["name"], "Иван Петров")
        self.assertEqual(len(items[0]["phones"]), 2)
        self.assertEqual(items[1]["tags"], ["инвестор", "vip"])
        for it in items:
            crm.upsert_contact(self.db, it, ["телефон"])
        rows = crm.select_contacts(self.db, ["инвестор"])
        self.assertEqual([r["phone"] for r in rows], ["+79031112233"])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM contacts").fetchone()[0], 2)

    def export(self, path, **kw):
        args = crm.build_parser().parse_args(["export", str(path)])
        for k, v in kw.items():
            setattr(args, k, v)
        crm.cmd_export(self.db, args)

    def test_export_edit_import_roundtrip(self):
        for it in crm.parse_vcf(VCF):
            crm.upsert_contact(self.db, it)
        path = self.dir / "c.csv"
        self.export(path)
        import csv
        rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines(),
                               delimiter=";"))
        self.assertEqual(rows[0][:3], ["id", "Телефон", "Имя в телефоне"])
        for r in rows:
            if r[1] == "+79161234567":
                r[7], r[8], r[12] = "стройка", "покупатель, москва", "да"
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, delimiter=";").writerows(rows)
        crm.cmd_import(self.db, type("A", (), {"file": str(path), "tag": ""}))
        r = self.db.execute("SELECT * FROM contacts WHERE phone='+79161234567'").fetchone()
        self.assertEqual(r["tags"], "покупатель,москва")
        self.assertEqual(r["sphere"], "стройка")
        self.assertEqual(r["consent"], 1)

    def test_xlsx_roundtrip_and_overlap(self):
        for it in crm.parse_vcf(VCF):
            it["in_phonebook"] = 1
            crm.upsert_contact(self.db, it)
        crm.upsert_contact(self.db, {"phones": ["+79161234567"], "name": "Ivan",
                                     "tags": ["чат-wlc"], "add_only": True})
        crm.upsert_contact(self.db, {"phones": ["+79990001122"], "name": "Olga",
                                     "tags": ["чат-premium-wlc"], "add_only": True})
        path = self.dir / "c.xlsx"
        self.export(path)
        import openpyxl
        ws = openpyxl.load_workbook(path).active
        rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["+79161234567"][10], "телефон + WLC")
        self.assertEqual(rows["+79161234567"][11], "да")
        self.assertIsNone(rows["+79990001122"][11])  # только в чате
        ws.cell(2, 8).value = "IT"
        openpyxl.load_workbook(path)  # файл читается
        items = crm.parse_table(path)
        self.assertEqual(len(items), 3)
        self.assertIn("чат-wlc", [t for i in items for t in i["tags"]])

    def test_classify(self):
        import classify
        self.assertEqual(classify.guess_gender("", "Галина Вайбер Зал"), "ж")
        self.assertEqual(classify.guess_gender("", "Кран Сергей"), "м")
        self.assertEqual(classify.guess_gender("", "Polina"), "ж")
        self.assertEqual(classify.guess_gender("", "Петрова р-р"), "ж")
        self.assertEqual(classify.guess_gender("", "Иванов Покупатель"), "м")
        self.assertEqual(classify.guess_gender("", "2ой Покупатель Снт Урал"), "")
        self.assertEqual(classify.region_of("+66806757116"), ("Таиланд", ""))
        self.assertEqual(classify.region_of("+77715017475"), ("Казахстан", ""))
        reg = self.dir / "numbering"
        reg.mkdir()
        (reg / "DEF-9xx.csv").write_text(
            "АВС/ DEF;От;До;Емкость;Оператор;Регион;ИНН\n"
            "916;0000000;9999999;10000000;ПАО \"МТС\";г. Москва и Московская область;1\n",
            encoding="utf-8")
        num = classify.Numbering(reg)
        self.assertEqual(classify.region_of("+79161234567", num),
                         ("г. Москва и Московская область", 'ПАО "МТС"'))
        rules = [("агент", ["р-р", "риелтор"]), ("wlc", ["wlc"]),
                 ("премиум", ["#чат-premium-wlc"])]
        self.assertEqual(classify.match_rules(rules, "Анна р-р", "", []), ["агент"])
        self.assertEqual(classify.match_rules(rules, "Рома Wlc", "из чата WLC",
                                              ["чат-premium-wlc"]), ["wlc", "премиум"])
        self.assertEqual(classify.match_rules(rules, "Рома", "из чата WLC", []), [])
        self.assertEqual(classify.match_rules([("агент", ["агент"])],
                                              "ALIMAR турагент", "агентство", []), [])

    def test_chat_list_only_adds_to_existing_contact(self):
        for it in crm.parse_vcf(VCF):
            crm.upsert_contact(self.db, it, ["клиент"])
        self.db.execute("UPDATE contacts SET consent=1")
        path = self.dir / "chat.csv"
        path.write_text("id;name;call_name;phone;telegram;tags;consent;stopped;notes\n"
                        ";Ivan;;+79161234567;;чат-1;;;из чата\n", encoding="utf-8-sig")
        crm.cmd_import(self.db, type("A", (), {"file": str(path), "tag": ""}))
        r = self.db.execute("SELECT * FROM contacts WHERE phone='+79161234567'").fetchone()
        self.assertEqual(r["name"], "Иван Петров")
        self.assertEqual(r["tags"], "клиент,чат-1")
        self.assertEqual(r["consent"], 1)
        self.assertIn("из чата", r["notes"])

    def test_render(self):
        row = {"name": "Иван Петров", "phone": "+7", "notes": ""}
        self.assertIn("Иван!", crm.render("{hello}, {first_name}! {unknown}", row))
        self.assertTrue(crm.render("{hello}, {first_name}!", {"name": "", "phone": "",
                                                              "notes": ""}).endswith("!"))

    def test_send_whatsapp_links_respects_consent_and_stop(self):
        crm.DATA_DIR = self.dir
        crm.CONFIG_PATH = self.dir / "none.ini"
        for it in crm.parse_vcf(VCF):
            crm.upsert_contact(self.db, it, ["клиент"])
        self.db.execute("UPDATE contacts SET consent=1")
        self.db.execute("UPDATE contacts SET stopped=1 WHERE phone='+79031112233'")
        tpl = self.dir / "t.txt"
        tpl.write_text("{hello}, {first_name}! Есть объект.", encoding="utf-8")
        args = crm.build_parser().parse_args(
            ["send", "--name", "Тест", "--template", str(tpl), "--tags", "клиент",
             "--channels", "wa", "--send"])
        crm.cmd_send(self.db, args)
        msgs = self.db.execute("SELECT * FROM messages").fetchall()
        self.assertEqual(len(msgs), 1)
        self.assertEqual(msgs[0]["status"], "link")
        page = next(self.dir.glob("wa_links_*.html")).read_text(encoding="utf-8")
        self.assertIn("https://wa.me/79161234567?text=", page)
        # повторный запуск не шлёт тем же людям
        crm.cmd_send(self.db, args)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM messages").fetchone()[0], 1)


if __name__ == "__main__":
    unittest.main()
