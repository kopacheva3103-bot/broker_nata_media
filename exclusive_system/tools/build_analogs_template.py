"""Шаблон «Анализ аналогов (конкуренты)»: сбор объявлений с ЦИАН/Авито + корректировки + итоговая цена.
Запуск: python3 tools/build_analogs_template.py → templates/analogs_template.xlsx (загружается в 02_ШАБЛОНЫ как Google Таблица)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(__file__), '..', 'templates', 'analogs_template.xlsx')
wb = Workbook()
HEAD = PatternFill('solid', fgColor='1F3864'); SUB = PatternFill('solid', fgColor='D9E2F3')
INP = PatternFill('solid', fgColor='FFF2CC'); CALC = PatternFill('solid', fgColor='EDEDED'); RES = PatternFill('solid', fgColor='E2EFDA')
W = Font(color='FFFFFF', bold=True); B = Font(bold=True); G = Font(color='595959', italic=True, size=9)
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

# ─────────────── Параметры ───────────────
P = wb.active; P.title = 'Параметры'
P['A1'] = 'ПАРАМЕТРЫ КОРРЕКТИРОВОК — жёлтые ячейки можно менять (ориентиры, утверждает брокер)'; P['A1'].font = B
rows = [
    ('Скидка на торг, % (продажа: жильё 3–5, коммерция 5–10, загородные 7–12; аренда 3–5)', 0.05, '0%'),
    ('Коэффициент торможения на площадь (коммерция −0,15; квартиры −0,10; дома −0,20)', -0.15, '0.00'),
    ('Разница площадей без корректировки, % (меньше — не корректируем)', 0.10, '0%'),
    ('Предельная общая корректировка аналога, % (больше — аналог не учитываем)', 0.30, '0%'),
]
for i, (label, v, fmt) in enumerate(rows, start=3):
    P.cell(i, 1, label).alignment = WRAP
    c = P.cell(i, 2, v); c.fill = INP; c.number_format = fmt; c.border = BOX
P['A8'] = 'РЕМОНТ (ПРОДАЖА) — стоимость ремонта от нуля, БЕЗ мебели (класс объекта — в блоке «Наш объект»)'; P['A8'].font = B
P['A9'], P['B9'], P['C9'] = 'Уровень', 'Комфорт, ₽ за м²', 'Премиум, ₽ за м²'
for c in ('A9', 'B9', 'C9'): P[c].fill = SUB; P[c].font = B
rep = [('Без отделки', 0, 0), ('С ремонтом (без мебели)', 100000, 250000)]
for i, (lvl, cm, pr) in enumerate(rep, start=10):
    P.cell(i, 1, lvl)
    for j, v in ((2, cm), (3, pr)):
        c = P.cell(i, j, v); c.fill = INP; c.number_format = '# ##0'; c.border = BOX
P['A12'] = 'Мебель и техника в корректировку ремонта не входят — если у аналога есть мебель, это ручная корректировка «Прочее».'; P['A12'].font = G
P['A13'] = 'Аренда: автоматической поправки на ремонт нет — ставку в месяц берём из объявлений ЦИАН, разницу в ремонте брокер ставит вручную («Прочее», %).'; P['A13'].font = G
P['A16'] = 'ВЕС ПО УРОВНЮ СОПОСТАВИМОСТИ'; P['A16'].font = B
for i, (lvl, w) in enumerate(((1, 1.0), (2, 0.7), (3, 0.4)), start=17):
    P.cell(i, 1, lvl); c = P.cell(i, 2, w); c.fill = INP; c.border = BOX
P['C17'] = '1 — прямой конкурент (та же локация, тип, площадь ±30 %)'
P['C18'] = '2 — расширенный поиск (соседний район / площадь ±50 % / снятые до 6 мес.)'
P['C19'] = '3 — дальний аналог (другая локация того же класса / снятые до 12 мес.)'
P.column_dimensions['A'].width = 70; P.column_dimensions['B'].width = 16; P.column_dimensions['C'].width = 60; P.column_dimensions['D'].width = 22

# ─────────────── Аналоги ───────────────
S = wb.create_sheet('Аналоги', 0)
S['A1'] = 'АНАЛИЗ АНАЛОГОВ (КОНКУРЕНТЫ) — объект:'; S['A1'].font = Font(bold=True, size=14)
S['A2'] = ('Жёлтое заполняет ассистент, серое считается само. Каждое объявление — строка + скриншот в папку «Аналитика» '
           '(объявления снимают). Цена аренды — ₽ в месяц за весь объект (ЦИАН для коммерции показывает ₽/м² в ГОД — пересчитать).')
S['A2'].font = G
S.merge_cells('A2:N2')

S['A4'] = 'НАШ ОБЪЕКТ'; S['A4'].font = W; S['A4'].fill = HEAD; S.merge_cells('A4:B4')
obj = [('Объект', ''), ('Адрес', ''), ('Сделка', 'Продажа'), ('Тип / назначение', ''), ('Площадь, м²', None),
       ('Ремонт', 'С ремонтом (без мебели)'), ('Этаж / вход', ''), ('Цена собственника, ₽', None), ('Дата анализа', ''),
       ('Класс объекта', 'Комфорт')]
for i, (k, v) in enumerate(obj, start=5):
    S.cell(i, 1, k).font = B
    c = S.cell(i, 2, v); c.fill = INP; c.border = BOX
S['B9'].number_format = '# ##0.0'; S['B12'].number_format = '# ##0'
# ссылки на параметры
AREA, DEAL, REPAIR, OWNER, KLASS = '$B$9', '$B$7', '$B$10', '$B$12', '$B$14'
TORG, ELAST, THR, MAXC = 'Параметры!$B$3', 'Параметры!$B$4', 'Параметры!$B$5', 'Параметры!$B$6'
REPTAB = 'Параметры!$A$10:$C$11'

S['D4'] = 'ИТОГ'; S['D4'].font = W; S['D4'].fill = HEAD; S.merge_cells('D4:G4')
F0, F1 = 21, 60  # строки аналогов
RCOL = 'IF(%s="Премиум",3,2)' % KLASS  # столбец стоимости ремонта: премиум / комфорт
rng = lambda col: '$%s$%d:$%s$%d' % (col, F0, col, F1)
use = rng('AD')      # 1 — аналог учитывается, 0 — нет (служебный столбец)
XU = rng('AC')       # скорректированная цена учтённых аналогов, иначе 0
res = [
    ('Аналогов учтено', '=SUM(%s)' % use, '0'),
    ('Средневзвешенная цена, ₽ за м²', '=IF(E5=0,"",SUMPRODUCT(%s,%s))' % (XU, rng('AB')), '# ##0'),
    ('Рыночная цена объекта, ₽', '=IFERROR(E6*%s,"")' % AREA, '# ##0'),
    ('Диапазон: от, ₽', '=IF(E5=0,"",_xlfn.MINIFS(%s,%s,1)*%s)' % (XU, use, AREA), '# ##0'),
    ('Диапазон: до, ₽', '=IF(E5=0,"",_xlfn.MAXIFS(%s,%s,1)*%s)' % (XU, use, AREA), '# ##0'),
    ('Коэффициент вариации (≤ 20 % — выборка однородная)',
     '=IF(E5<2,"",SQRT(SUMPRODUCT(%s,(%s-SUM(%s)/E5)^2)/(E5-1))/(SUM(%s)/E5))' % (use, XU, XU, XU), '0%'),
    ('Цена собственника к рынку', '=IFERROR(%s/E7-1,"")' % OWNER, '+0%;-0%;0%'),
    ('Средний срок экспозиции учтённых, дней', '=IFERROR(SUMPRODUCT(%s,%s)/SUMPRODUCT(%s,--(%s<>"")),"")' % (use, rng('F'), use, rng('F')), '0'),
    ('Вывод', '=IF(E5<3,"Мало аналогов — расширьте поиск (уровень 2–3)",IF(E10>0.2,"Выборка разнородная — пересмотрите аналоги и корректировки",IF(E11>0.1,"Цена собственника выше рынка более чем на 10 %",IF(E11<-0.1,"Цена собственника ниже рынка более чем на 10 %","Цена собственника в рынке"))))', '@'),
]
for i, (k, f, fmt) in enumerate(res, start=5):
    S.cell(i, 4, k).alignment = WRAP
    c = S.cell(i, 5, f); c.fill = RES; c.number_format = fmt; c.border = BOX; c.font = B
S.merge_cells('E13:N13')

hdr = [
    ('№', 5, None), ('Уровень (1/2/3)', 8, 'inp'), ('Ссылка на объявление', 28, 'inp'), ('Дата сбора', 11, 'inp'),
    ('Статус', 10, 'inp'), ('Дней в экспозиции', 9, 'inp'), ('Адрес / ЖК / посёлок', 26, 'inp'), ('Расстояние до нашего, км', 9, 'inp'),
    ('Тип / назначение', 14, 'inp'), ('Площадь, м²', 9, 'inp'), ('Участок, сот. (дома)', 9, 'inp'), ('Этаж / вход', 12, 'inp'),
    ('Ремонт', 20, 'inp'), ('Особенности: чем лучше (+) / хуже (−) нашего', 30, 'inp'), ('Цена предложения, ₽ (аренда — ₽/мес)', 15, 'inp'),
    ('Цена за м², ₽', 11, 'calc'), ('Торг, %', 8, 'calc'), ('Площадь, %', 8, 'calc'), ('Ремонт, ₽ за м²', 10, 'calc'),
    ('Местоположение, %', 10, 'inp'), ('Этаж / вход / тип, %', 10, 'inp'), ('Прочее (парковка, вид, участок, мебель), %', 12, 'inp'),
    ('Обоснование ручных корректировок', 30, 'inp'), ('Скорректированная цена за м², ₽', 13, 'calc'),
    ('Общая корректировка без ремонта, %', 11, 'calc'), ('Учитываем?', 14, 'calc'), ('Вес (сырой)', 8, 'calc'), ('Вес', 7, 'calc'),
    ('служ.: цена учтённого', 9, 'calc'), ('служ.: учтён', 6, 'calc'),
]
HR = F0 - 1
S.cell(HR - 1, 1, 'АНАЛОГИ').font = W; S.cell(HR - 1, 1).fill = HEAD
S.cell(HR - 1, 16, 'КОРРЕКТИРОВКИ (порядок: торг → площадь → ремонт → местоположение, этаж, прочее)').font = W
for col in range(16, 31): S.cell(HR - 1, col).fill = HEAD
S.column_dimensions.group('AC', 'AD', hidden=True)
for j, (name, w, kind) in enumerate(hdr, start=1):
    c = S.cell(HR, j, name); c.font = B; c.fill = SUB; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = BOX
    S.column_dimensions[c.column_letter].width = w
S.row_dimensions[HR].height = 60
for r in range(F0, F1 + 1):
    e = lambda f: '=IF($O%d="","",%s)' % (r, f)
    vals = {
        1: r - F0 + 1,
        16: e('O%d/J%d' % (r, r)),
        17: e('IF(E%d="Сделка",0,-%s)' % (r, TORG)),
        18: e('IF(ABS(J%d/%s-1)<=%s,0,(%s/J%d)^%s-1)' % (r, AREA, THR, AREA, r, ELAST)),
        19: e('IF(%s="Аренда",0,IFERROR(VLOOKUP(%s,%s,%s,FALSE)-VLOOKUP(M%d,%s,%s,FALSE),0))' % (DEAL, REPAIR, REPTAB, RCOL, r, REPTAB, RCOL)),
        24: e('(P%d*(1+Q%d)*(1+R%d)+S%d)*(1+N(T%d))*(1+N(U%d))*(1+N(V%d))' % (r, r, r, r, r, r, r)),
        25: e('ABS(Q%d)+ABS(R%d)+ABS(N(T%d))+ABS(N(U%d))+ABS(N(V%d))' % (r, r, r, r, r)),
        26: e('IF(Y%d>%s,"Нет — корректировки > предела","Да")' % (r, MAXC)),
        27: e('IF(Z%d="Да",IFERROR(VLOOKUP(B%d,Параметры!$A$17:$B$19,2,FALSE),0.4)/(Y%d+ABS(S%d/(P%d*(1+Q%d)))+0.1),0)' % (r, r, r, r, r, r)),
        28: '=IFERROR(AA%d/SUM($AA$%d:$AA$%d),0)' % (r, F0, F1),
        29: '=IF(Z%d="Да",X%d,0)' % (r, r),
        30: '=IF(Z%d="Да",1,0)' % r,
    }
    for j, (name, w, kind) in enumerate(hdr, start=1):
        c = S.cell(r, j, vals.get(j)); c.border = BOX
        c.fill = INP if kind == 'inp' else CALC if kind == 'calc' else PatternFill()
    for j, fmt in ((10, '# ##0.0'), (15, '# ##0'), (16, '# ##0'), (17, '0%'), (18, '+0.0%;-0.0%;0%'), (19, '+# ##0;-# ##0;0'),
                   (20, '+0%;-0%;0%'), (21, '+0%;-0%;0%'), (22, '+0%;-0%;0%'), (24, '# ##0'), (25, '0%'), (28, '0%'), (4, 'dd.mm.yyyy')):
        S.cell(r, j).number_format = fmt
S.freeze_panes = S.cell(F0, 4)

def dv(formula, ref):
    d = DataValidation(type='list', formula1=formula, allow_blank=True); S.add_data_validation(d); d.add(ref)
dv('"1,2,3"', 'B%d:B%d' % (F0, F1))
dv('"Активно,Снято,Сделка"', 'E%d:E%d' % (F0, F1))
dv('Параметры!$A$10:$A$11', 'M%d:M%d' % (F0, F1))
dv('Параметры!$A$10:$A$11', 'B10')
dv('"Комфорт,Премиум"', 'B14')
dv('"Продажа,Аренда"', 'B7')
S['E13'].alignment = WRAP
S.column_dimensions['D'].width = 34

# ─────────────── Как заполнять ───────────────
H = wb.create_sheet('Как заполнять')
steps = [
    'КАК ЗАПОЛНЯТЬ (подробно — инструкция «09 — Анализ конкурентов» в 03_ИНСТРУКЦИИ)',
    '1. Сделайте копию шаблона в папку объекта «Аналитика»: «Анализ аналогов — <объект> — дд.мм.гггг».',
    '2. Заполните блок «Наш объект» (сделка, площадь, ремонт, класс — обязательно). Ремонт (продажа): комфорт — 100 000 ₽/м², премиум — 250 000 ₽/м², без мебели. Аренда — ставка в месяц из ЦИАН, ремонт вручную.',
    '3. ЦИАН: те же тип и сделка, та же локация, площадь ±30 %. Сортировка по дате. Минимум 5 прямых аналогов (уровень 1).',
    '4. Если прямых меньше 5 — расширяйте по шагам: радиус → площадь ±50 % → снятые объявления до 6–12 мес. → другая локация того же класса. Уровень ставьте 2 или 3.',
    '5. По каждому объявлению: ссылка, скриншот, площадь, ремонт, этаж/вход, цена, дней в экспозиции, особенности (+/−).',
    '6. Корректировки на торг, площадь и ремонт считаются сами. Местоположение, этаж/вход и прочее ставит брокер (в % с обоснованием).',
    '7. Аналог с общей корректировкой больше 30 % (без учёта ремонта) не учитывается автоматически. Разница в ремонте снижает вес аналога.',
    '8. Итог (справа вверху): средневзвешенная цена за м², цена объекта, диапазон, однородность, сравнение с ценой собственника.',
    '9. Файл → в Claude с промптом «Анализ цены по аналогам» → ответ «Вставить стратегию из Claude».',
]
for i, t in enumerate(steps, start=1):
    H.cell(i, 1, t).alignment = WRAP
H['A1'].font = B; H.column_dimensions['A'].width = 130

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT)
