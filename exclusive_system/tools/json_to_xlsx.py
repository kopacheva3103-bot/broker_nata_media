import json, sys, datetime, re
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
src = json.load(open(sys.argv[1]))
order = ['01_ОБЪЕКТЫ','02_СТРАТЕГИЯ','03_ДЕЙСТВИЯ','04_ЛИДЫ_И_КОНВЕРСИИ','05_СТАТИСТИКА','06_ПЛАН_ФАКТ','07_ОТЧЕТ','08_СПРАВОЧНИКИ','09_ДЭШБОРД','10_НАСТРОЙКИ','11_КОНТРОЛЬ','12_ИСТОРИЯ','13_АРХИВ_ОТЧЕТОВ']
wb = Workbook(); wb.remove(wb.active)
for name in order:
    ws = wb.create_sheet(name)
    for a1, c in src['sheets'].get(name, {}).items():
        if 'f' in c: ws[a1] = c['f']
        elif 'd' in c: ws[a1] = datetime.datetime(*c['d'])
        else:
            v = c['v']
            if v == '' or v is None: continue
            ws[a1] = v
for n, t in src['names'].items():
    col, row = re.match(r'([A-Z]+)(\d+)', t['a1']).groups()
    wb.defined_names[n] = DefinedName(n, attr_text="'%s'!$%s$%s" % (t['sheet'], col, row))
wb.save(sys.argv[2])
print('saved')
